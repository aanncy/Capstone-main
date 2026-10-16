import re
import asyncio
import json
import xml.etree.ElementTree as ET
from typing import get_args

import pytest

from fastapi.testclient import TestClient

from app.convert import to_drawio_xml
from app.main import app, get_provider
from app.pipeline import PipelineError, generate_diagram, parse_json
from app.providers import ProviderChain, ProviderError
from app.schema import Diagram, NodeType

GOOD = {
    "title": "Payment request",
    "lanes": [{"id": "maker", "name": "Maker"}, {"id": "approver", "name": "Approver"}],
    "nodes": [
        {"id": "n1", "type": "start", "label": "Request received", "lane": "maker", "source_ref": "P1", "support": "stated"},
        {"id": "n2", "type": "task", "label": "Prepare documents", "lane": "maker", "source_ref": "P1", "support": "stated"},
        {"id": "n3", "type": "decision", "label": "Within limit?", "lane": "approver", "source_ref": "P1", "support": "implied"},
        {"id": "n4", "type": "end", "label": "Payment released", "lane": "approver", "source_ref": "P1", "support": "stated"},
    ],
    "edges": [
        {"id": "e1", "from": "n1", "to": "n2", "label": ""},
        {"id": "e2", "from": "n2", "to": "n3", "label": ""},
        {"id": "e3", "from": "n3", "to": "n4", "label": "Yes"},
        {"id": "e4", "from": "n3", "to": "n2", "label": "No"},
    ],
}


@pytest.mark.parametrize("kind", get_args(NodeType))
def test_all_node_symbols_convert_with_connected_geometry(kind):
    from app.convert import SHAPES, STYLES
    data = json.loads(json.dumps(GOOD))
    data["nodes"][1]["type"] = kind
    root = ET.fromstring(to_drawio_xml(Diagram.model_validate(data)))
    cell = root.find(".//object[@id='n2']/mxCell")
    # the type's own style, plus margins that keep the text inside the outline
    assert cell is not None and cell.get("style").startswith(STYLES[kind]) and "spacingLeft=" in cell.get("style")
    geometry = cell.find("mxGeometry")
    assert float(geometry.get("width")) >= SHAPES[kind][0]
    assert float(geometry.get("height")) >= SHAPES[kind][1]
    if kind == "connector":
        assert geometry.get("width") == geometry.get("height")
    assert root.find(".//mxCell[@source='n1'][@target='n2']") is not None
    assert root.find(".//mxCell[@source='n2'][@target='n3']") is not None


class Fake:
    name = "fake"

    def __init__(self, *replies):
        self.replies, self.calls, self.attachments = list(replies), 0, []

    async def complete(self, system, user, attachments=()):
        self.calls += 1
        self.attachments.append(list(attachments))
        return self.replies.pop(0)


def test_parse_json_handles_fences_and_chatter():
    assert parse_json("Sure!\n```json\n" + json.dumps(GOOD) + "\n```")["title"] == "Payment request"


def test_valid_reply_needs_one_call():
    p = Fake(json.dumps(GOOD))
    d = asyncio.run(generate_diagram("text", p))
    assert p.calls == 1 and len(d.nodes) == 4


def test_bad_reply_is_fixed_automatically():
    bad = json.loads(json.dumps(GOOD))
    bad["edges"] = [e for e in bad["edges"] if e["label"] != "No"]  # decision with one exit
    p = Fake(json.dumps(bad), json.dumps(GOOD))
    d = asyncio.run(generate_diagram("text", p))
    assert p.calls == 2 and d.title == "Payment request"


def test_gives_up_after_retries():
    p = Fake("not json", "still not json", "nope")
    try:
        asyncio.run(generate_diagram("text", p))
        assert False, "should have raised"
    except PipelineError as e:
        assert "3 tries" in str(e) and p.calls == 3


def test_xml_is_well_formed_and_keeps_custom_fields():
    from app.schema import Diagram
    lanes = json.loads(json.dumps(GOOD))
    lanes["lanes"][1]["support"] = "guess"
    lanes["edges"][3].update(support="implied", sources=["SOP.pdf", "call.txt"])
    root = ET.fromstring(to_drawio_xml(Diagram.model_validate(lanes)))
    objs = {o.get("id"): o for o in root.iter("object")}
    assert set(objs) == {"0", "n1", "n2", "n3", "n4", "lane_maker", "lane_approver", "e1", "e2", "e3", "e4"}
    assert objs["n3"].get("support") == "implied" and objs["lane_approver"].get("support") == "guess"
    # the unsure arrow (n3 -> n2, "No") is dashed, the others are solid
    assert objs["e4"].get("support") == "implied" and objs["e4"].get("label") == "No"
    assert objs["e4"].get("sources") == "SOP.pdf; call.txt"
    assert "dashed=1" in objs["e4"].find("mxCell").get("style")
    assert "dashed=1" not in objs["e1"].find("mxCell").get("style")
    ids = [e.get("id") for e in root.iter() if e.get("id")]
    assert len(ids) == len(set(ids))


def test_chain_falls_back_to_next_provider():
    class Down:
        name = "down"

        async def complete(self, s, u, a=()):
            raise ProviderError("rate limit")

    assert asyncio.run(ProviderChain([Down(), Fake("ok")]).complete("s", "u")) == "ok"


def test_endpoint_returns_xml_and_clear_errors():
    app.dependency_overrides[get_provider] = lambda: Fake(json.dumps(GOOD))
    c = TestClient(app)
    r = c.post("/api/generate", data={"text": "A process"})
    assert r.status_code == 200 and r.json()["xml"].startswith("<mxfile")
    r = c.post("/api/generate", data={"text": "  "})
    assert r.status_code == 400 and "Add a description" in r.json()["detail"]
    r = c.post("/api/generate", files={"file": ("notes.xls", b"x")})
    assert r.status_code == 400 and "Unsupported file type" in r.json()["detail"]
    app.dependency_overrides.clear()


def test_files_are_sent_to_the_model_on_every_try():
    bad = json.loads(json.dumps(GOOD))
    bad["edges"] = bad["edges"][:-1]
    fake = Fake(json.dumps(bad), json.dumps(GOOD))
    app.dependency_overrides[get_provider] = lambda: fake
    r = TestClient(app).post("/api/generate", files={"file": ("flow.pdf", b"%PDF-1.4 scan")})
    app.dependency_overrides.clear()
    assert r.status_code == 200 and fake.calls == 2
    assert all([(a.name, a.mime, a.data) for a in call] == [("flow.pdf", "application/pdf", b"%PDF-1.4 scan")]
               for call in fake.attachments)


def test_word_file_sends_text_and_pictures():
    import base64
    import io
    from docx import Document
    from app.extract import _docx
    png = base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNk+M9QDwADhgGAWjR9awAAAABJRU5ErkJggg==")
    doc = Document()
    doc.add_paragraph("Maker prepares the request.")
    doc.add_picture(io.BytesIO(png))
    buf = io.BytesIO()
    doc.save(buf)
    out = _docx("p.docx", buf.getvalue())
    assert [a.mime for a in out] == ["text/plain", "image/png"] and b"Maker" in out[0].data and out[1].data == png


def test_text_only_providers_take_text_files_but_refuse_pdfs():
    from app.extract import Attachment
    from app.providers import text_only
    assert "hello" in text_only("u", [Attachment("a.txt", "text/plain", b"hello")], "X")
    try:
        text_only("u", [Attachment("a.pdf", "application/pdf", b"%PDF")], "X")
        assert False, "should have raised"
    except ProviderError as e:
        assert "PROVIDERS=gemini" in str(e)


def test_runs_are_saved_side_by_side_per_input(tmp_path, monkeypatch):
    from app.extract import Attachment
    from app.history import save_run
    from app.schema import Diagram
    monkeypatch.setenv("SAVE_OUTPUTS", "true")
    monkeypatch.setenv("OUTPUT_DIR", str(tmp_path))
    d = Diagram.model_validate(GOOD)
    pdf = [Attachment("flow.pdf", "application/pdf", b"%PDF")]
    tries = [{"attempt": 1, "reply": "oops", "problems": ["not JSON"], "thoughts": "Found two roles."},
             {"attempt": 2, "reply": "```json\n" + json.dumps(GOOD) + "\n```", "problems": []}]
    a = save_run("", "flow.pdf", pdf, {"model": "gemini/flash"}, tries, d, to_drawio_xml(d))
    b = save_run("", "flow.pdf", pdf, {"model": "github_models/openai/gpt-4.1"}, tries[:1], error="gave up")
    assert a.parent == b.parent and a.parent.name.startswith("flow-") and b.name.endswith("gpt-4.1_FAILED.json")
    stem = a.with_suffix("")
    assert stem.with_suffix(".drawio").exists() and not b.with_suffix(".drawio").exists()
    assert (a.parent / f"{stem.name}_attempt1.txt").read_text() == "oops"
    assert (a.parent / f"{stem.name}_attempt1_thoughts.md").read_text() == "Found two roles."
    assert not (a.parent / f"{stem.name}_attempt2_thoughts.md").exists()
    assert json.loads((a.parent / f"{stem.name}_attempt2.json").read_text())["title"] == "Payment request"
    meta = json.loads(a.read_text())
    assert meta["run"]["attempts"][0]["problems"] == ["not JSON"] and meta["diagram"]["title"] == "Payment request"



def test_history_lists_runs_with_a_diagram_and_reopens_them(tmp_path, monkeypatch):
    from app.extract import Attachment
    from app.history import save_run
    from app.schema import Diagram
    monkeypatch.setenv("SAVE_OUTPUTS", "true")
    monkeypatch.setenv("OUTPUT_DIR", str(tmp_path))
    d = Diagram.model_validate(GOOD)
    word = [Attachment("sop.docx", "text/plain", b"steps"), Attachment("sop.docx picture 3", "image/png", b"png")]
    save_run("notes", "sop.docx", word, {"model": "gemini/flash"}, (), d, to_drawio_xml(d))  # saved before "files"
    save_run("", "a +1 more", [], {"files": ["a.pdf", "b.txt"], "model": "gemini/pro"}, (), error="gave up")
    (tmp_path / "secret.drawio").write_text("<mxfile/>")  # outside any run folder: never served

    client = TestClient(app)
    runs = client.get("/api/history").json()["runs"]
    assert len(runs) == 1  # the failed run has no diagram to open
    run = runs[0]
    assert run["documents"] == ["sop.docx"] and run["typed"] and run["model"] == "gemini/flash"
    assert run["kind"] == "generate" and run["title"] == "Payment request" and run["time"]
    assert "<mxfile" in client.get(f"/api/history/{run['id']}").json()["xml"]
    assert client.get("/api/history/nope/nope").status_code == 404
    assert client.get("/api/history/..%2F/secret").status_code == 404
    assert client.get("/api/history/./secret").status_code == 404

def test_gemini_thoughts_are_kept_out_of_the_answer(monkeypatch):
    import httpx
    from app.providers import GeminiProvider
    sent = {}

    async def post(self, url, headers=None, json=None):
        sent.update(json)
        parts = [{"text": "Lanes: Maker, Approver.", "thought": True}, {"text": '{"title": "x"}'}]
        return httpx.Response(200, json={"candidates": [{"content": {"parts": parts}}]})
    monkeypatch.setattr(httpx.AsyncClient, "post", post)
    chain = ProviderChain([GeminiProvider("key", "gemini-test", thoughts=True)])
    assert asyncio.run(chain.complete("s", "u")) == '{"title": "x"}'
    assert chain.thoughts == "Lanes: Maker, Approver."
    assert sent["generationConfig"]["thinkingConfig"] == {"includeThoughts": True}


def test_drawio_xml_reads_back_into_the_same_diagram():
    import base64, zlib
    from urllib.parse import quote
    from app.read_drawio import from_drawio_xml
    from app.schema import Diagram
    d = Diagram.model_validate(GOOD)
    xml = to_drawio_xml(d)
    assert from_drawio_xml(xml).model_dump() == d.model_dump()
    # draw.io may save the page compressed
    model = re.search(r"<mxGraphModel.*</mxGraphModel>", xml).group(0)
    deflate = zlib.compressobj(9, zlib.DEFLATED, -15)
    packed = base64.b64encode(deflate.compress(quote(model).encode()) + deflate.flush()).decode()
    compressed = f'<mxfile><diagram id="d1" name="Payment request">{packed}</diagram></mxfile>'
    assert from_drawio_xml(compressed).model_dump() == d.model_dump()


def test_shapes_drawn_by_hand_are_read_with_a_type_and_lane():
    from app.read_drawio import from_drawio_xml
    from app.schema import Diagram
    xml = to_drawio_xml(Diagram.model_validate(GOOD)).replace("</root>", (
        '<mxCell id="x1" value="Call &lt;b&gt;customer&lt;/b&gt;" style="rhombus;whiteSpace=wrap;" vertex="1" parent="lane_maker">'
        '<mxGeometry x="10" y="10" width="80" height="80" as="geometry"/></mxCell>'
        '<mxCell id="x2" value="" style="endArrow=classic;" edge="1" parent="1" source="n2" target="x1">'
        '<mxGeometry relative="1" as="geometry"/></mxCell>'
        '<mxCell id="x3" value="Yes" style="edgeLabel;" vertex="1" connectable="0" parent="x2">'
        '<mxGeometry relative="1" as="geometry"/></mxCell></root>'))
    d = from_drawio_xml(xml)
    x1 = next(n for n in d.nodes if n.id == "x1")
    assert (x1.type, x1.lane, x1.label, x1.source_ref) == ("decision", "maker", "Call customer", "Added in the editor")
    assert any(e.source == "n2" and e.target == "x1" and e.label == "Yes" for e in d.edges)
    assert len(d.nodes) == 5  # the arrow's label cell is not a step


def test_amend_sends_the_current_diagram_and_returns_the_new_one():
    changed = json.loads(json.dumps(GOOD))
    changed["nodes"].insert(2, {"id": "n5", "type": "task", "label": "Check compliance", "lane": "approver"})
    changed["edges"][1] = {"from": "n2", "to": "n5", "label": ""}
    changed["edges"].append({"from": "n5", "to": "n3", "label": ""})
    fake = Fake(json.dumps(changed))
    seen = []
    real = fake.complete

    async def spy(system, user, attachments=()):
        seen.append((system, user))
        return await real(system, user, attachments)
    fake.complete = spy
    app.dependency_overrides[get_provider] = lambda: fake
    from app.schema import Diagram
    c = TestClient(app)
    r = c.post("/api/amend", json={"xml": to_drawio_xml(Diagram.model_validate(GOOD)), "instruction": "Add a compliance check"})
    assert r.status_code == 200 and 'id="n5"' in r.json()["xml"]
    system, user = seen[0]
    assert "Amending an existing diagram" in system
    assert "Add a compliance check" in user and '"Prepare documents"' in user
    assert c.post("/api/amend", json={"xml": "<mxfile/>", "instruction": "x"}).status_code == 400
    assert c.post("/api/amend", json={"xml": "x", "instruction": " "}).status_code == 400
    app.dependency_overrides.clear()


def test_lanes_are_tinted_and_named_along_their_length():
    from app.schema import Diagram
    long = json.loads(json.dumps(GOOD))  # a long flow, so the lane name has to repeat
    long["nodes"] = [GOOD["nodes"][0]] + [
        {"id": f"t{i}", "type": "task", "label": f"Step {i}", "lane": "maker" if i % 2 else "approver"} for i in range(12)
    ] + [GOOD["nodes"][3]]
    ids = [n["id"] for n in long["nodes"]]
    long["edges"] = [{"from": a, "to": b} for a, b in zip(ids, ids[1:])]
    root = ET.fromstring(to_drawio_xml(Diagram.model_validate(long)))
    cells = {c.get("id"): c for c in root.iter("mxCell") if c.get("id")}
    lanes = {o.get("id"): o.find("mxCell").get("style") for o in root.iter("object") if o.get("id", "").startswith("lane_")}
    assert len({s.split("swimlaneFillColor=")[1].split(";")[0] for s in lanes.values()}) == 2
    tags = [c for c in cells.values() if c.get("id").startswith("tag_maker_")]
    assert len(tags) >= 2 and all(t.get("value") == "Maker" for t in tags)
    steps = [o.find("mxCell") for o in root.iter("object") if o.get("node_type")]
    for t in tags:  # a repeated name never sits on a step
        g = t.find("mxGeometry")
        tx, tw = float(g.get("x")), float(g.get("width"))
        for s in steps:
            sg = s.find("mxGeometry")
            if s.get("parent") == "lane_maker":
                sx, sy = float(sg.get("x")), float(sg.get("y"))
                assert sx + float(sg.get("width")) < tx or sx > tx + tw or sy > float(g.get("y")) + float(g.get("height"))



def test_labels_fit_inside_their_shape():
    import textwrap
    from app.convert import CHAR_W, LINE_H, SHAPES, _size
    from app.schema import Node
    labels = ["Approve", "Perform maker-checker/verification review of the instruction",
              "Is the high-risk data condition detected for this outward payment?",
              "Supercalifragilisticexpialidocious verification"]
    for kind in set(get_args(NodeType)) - {"parallel"}:  # a parallel bar hides its label
        for label in labels:
            w, h = _size(Node(id="x", type=kind, label=label, lane="l"))
            left, right, top, bottom = SHAPES[kind][3]
            chars = int((w * (1 - left - right) - 8) // CHAR_W)
            lines = textwrap.wrap(label, chars)
            assert len(lines) * LINE_H + 8 <= h * (1 - top - bottom) + 1, (kind, label, w, h)
            if w < SHAPES[kind][2]:  # below the widest allowed, a long word is never split
                assert all(len(word) <= chars for word in label.split()), (kind, label, w)


def test_several_files_are_collated_in_one_request(monkeypatch):
    fake = Fake(json.dumps(GOOD))
    app.dependency_overrides[get_provider] = lambda: fake
    c = TestClient(app)
    sources = [("files", ("SOP.pdf", b"%PDF-1.4 sop")), ("files", ("transcript.txt", b"Mary: the team lead approves now")),
               ("files", ("sketch.png", b"\x89PNG rough"))]
    r = c.post("/api/generate", data={"text": "Payments process"}, files=sources)
    assert r.status_code == 200 and fake.calls == 1
    assert [a.name for a in fake.attachments[0]] == ["SOP.pdf", "transcript.txt", "sketch.png"]
    # too many files, or files that together exceed the request limit, are refused with a clear message
    r = c.post("/api/generate", files=[("files", (f"f{i}.txt", b"x")) for i in range(11)])
    assert r.status_code == 400 and "at most 10" in r.json()["detail"]
    big = b"%PDF" + b"0" * (8 * 1024 * 1024)
    r = c.post("/api/generate", files=[("files", ("a.pdf", big)), ("files", ("b.pdf", big))])
    assert r.status_code == 413 and "add up to" in r.json()["detail"]
    # a bad file names which one
    r = c.post("/api/generate", files=[("files", ("ok.txt", b"x")), ("files", ("notes.xls", b"x"))])
    assert r.status_code == 400 and r.json()["detail"].startswith("notes.xls:")
    app.dependency_overrides.clear()


def test_all_inputs_reach_generation_and_fix_rounds():
    bad = json.loads(json.dumps(GOOD))
    bad["edges"] = bad["edges"][:-1]
    fake = Fake(json.dumps(bad), json.dumps(GOOD))
    prompts = []
    complete = fake.complete

    async def capture(system, user, attachments=()):
        prompts.append(user)
        return await complete(system, user, attachments)

    fake.complete = capture
    app.dependency_overrides[get_provider] = lambda: fake
    sources = [("files", ("procedure.pdf", b"%PDF procedure")),
               ("files", ("changes.txt", b"Add compliance review")),
               ("files", ("flow.png", b"image source"))]
    try:
        response = TestClient(app).post("/api/generate", data={"text": "Include reconciliation"}, files=sources)
        assert response.status_code == 200 and fake.calls == 2
        for prompt, attachments in zip(prompts, fake.attachments):
            assert "Include reconciliation" in prompt
            assert [a.name for a in attachments] == ["procedure.pdf", "changes.txt", "flow.png"]
            assert [a.data for a in attachments] == [b"%PDF procedure", b"Add compliance review", b"image source"]
            assert all(name in prompt for name in ["procedure.pdf", "changes.txt", "flow.png"])
    finally:
        app.dependency_overrides.clear()


def test_text_limit_depends_on_the_provider(monkeypatch):
    from app.main import max_chars
    monkeypatch.delenv("MAX_INPUT_CHARS", raising=False)
    monkeypatch.setenv("PROVIDERS", "gemini,github_models")
    assert max_chars() == 300_000  # a long meeting transcript fits
    monkeypatch.setenv("PROVIDERS", "github_models")
    assert max_chars() == 15_000
    monkeypatch.setenv("MAX_INPUT_CHARS", "50000")
    assert max_chars() == 50_000


def test_gemini_is_told_which_file_is_which(monkeypatch):
    import httpx
    from app.extract import Attachment
    from app.providers import GeminiProvider
    sent = {}

    async def post(self, url, headers=None, json=None):
        sent.update(json)
        return httpx.Response(200, json={"candidates": [{"content": {"parts": [{"text": "{}"}]}}]})
    monkeypatch.setattr(httpx.AsyncClient, "post", post)
    files = [Attachment("SOP.pdf", "application/pdf", b"%PDF"), Attachment("call.txt", "text/plain", b"hi")]
    asyncio.run(GeminiProvider("key", "m").complete("s", "u", files))
    parts = sent["contents"][0]["parts"]
    assert [p.get("text") for p in parts[:4:2]] == ["SOURCE 1: SOP.pdf", "SOURCE 2: call.txt"]
    assert parts[1]["inlineData"]["mimeType"] == "application/pdf" and parts[-1] == {"text": "u"}


def test_support_categories_replace_confidence_numbers():
    from app.read_drawio import from_drawio_xml
    from app.schema import Node
    # diagrams saved before categories had numbers; 1 meant a person had reviewed the item
    as_support = [Node(id="x", type="task", label="a", lane="l", confidence=c).support for c in (1, .95, .8, .5, .2)]
    assert as_support == ["reviewed", "stated", "derived", "implied", "guess"]
    old = to_drawio_xml(Diagram.model_validate(GOOD)).replace('support="implied"', 'confidence="0.55"')
    assert next(n for n in from_drawio_xml(old).nodes if n.id == "n3").support == "implied"


def test_conflicts_questions_and_skipped_sections_survive_the_editor():
    from app.read_drawio import from_drawio_xml
    from app.schema import check_logic
    data = json.loads(json.dumps(GOOD))
    data["conflicts"] = [{"affects": ["n3", "e4"], "description": "SOP: Checker; call: Team Lead",
                          "resolution": "Drew Team Lead: the call says it changed"}]
    data["open_questions"] = [{"question": "Who approves over 200k?", "affects": ["lane_x"], "source_ref": "call.txt"}]
    data["skipped_sections"] = [{"source": "SOP.pdf", "section": "1. Document control", "reason": "not a step"}]
    d = Diagram.model_validate(data)
    assert check_logic(d) == ["Open question 1 refers to ids that do not exist: lane_x."]
    data["open_questions"][0]["affects"] = ["approver"]
    d = Diagram.model_validate(data)
    assert check_logic(d) == []
    back = from_drawio_xml(to_drawio_xml(d))
    assert back.model_dump() == d.model_dump()


def test_new_diagrams_use_the_short_list_of_node_types_but_amends_keep_others():
    from app.pipeline import amend_diagram
    special = json.loads(json.dumps(GOOD))
    special["nodes"][1]["type"] = "manual_input"
    p = Fake(json.dumps(special), json.dumps(GOOD))
    asyncio.run(generate_diagram("text", p))
    assert p.calls == 2  # manual_input sent it back for a fix
    p = Fake(json.dumps(special))
    assert asyncio.run(amend_diagram(Diagram.model_validate(special), "rename", p)).nodes[1].type == "manual_input"


def test_reviews_survive_an_amend_and_the_ai_cannot_review():
    from app.pipeline import amend_diagram
    before = json.loads(json.dumps(GOOD))
    before["nodes"][2]["support"] = "reviewed"   # n3 checked by a person
    after = json.loads(json.dumps(GOOD))          # the AI dropped n3's "reviewed" and claims n1 is reviewed
    after["nodes"][0]["support"] = "reviewed"
    after["nodes"][3]["label"] = "Payment sent"   # and the change renamed n4
    d = asyncio.run(amend_diagram(Diagram.model_validate(before), "rename the end", Fake(json.dumps(after))))
    support = {n.id: n.support for n in d.nodes}
    assert support["n3"] == "reviewed" and support["n1"] == "stated"


def test_structured_output_schema_reaches_gemini_with_a_fallback(monkeypatch):
    import httpx
    from app.providers import GeminiProvider
    from app.prompts import response_schema
    bodies = []

    async def post(self, url, headers=None, json=None):
        bodies.append(__import__("copy").deepcopy(json))
        if "responseJsonSchema" in json["generationConfig"]:
            return httpx.Response(400, json={"error": {"message": "schema not supported"}})
        return httpx.Response(200, json={"candidates": [{"content": {"parts": [{"text": "{}"}]}}]})
    monkeypatch.setattr(httpx.AsyncClient, "post", post)
    chain = ProviderChain([GeminiProvider("key", "m")])
    asyncio.run(chain.complete("s", "u", (), schema=response_schema()))
    assert len(bodies) == 2 and "responseJsonSchema" not in bodies[1]["generationConfig"]
    node_types = bodies[0]["generationConfig"]["responseJsonSchema"]["properties"]["nodes"]["items"]["properties"]["type"]
    assert node_types["enum"] == ["start", "end", "task", "decision", "parallel", "subprocess", "document", "database"]


def test_a_parallel_bar_splits_or_joins():
    from app.schema import check_logic
    data = json.loads(json.dumps(GOOD))
    data["nodes"][1].update(type="parallel", label="Run checks in parallel")
    data["edges"][3]["to"] = "n4"  # the "No" loop no longer comes back to n2: one arrow in, one out
    errs = check_logic(Diagram.model_validate(data))
    assert errs == ["Parallel step 'n2' must split the flow (2+ outgoing arrows) or join it (2+ incoming)."]
    data["edges"] += [{"from": "n1", "to": "n2"}]  # now two arrows come in: a join
    assert check_logic(Diagram.model_validate(data)) == []
    cell = ET.fromstring(to_drawio_xml(Diagram.model_validate(data))).find(".//object[@id='n2']")
    assert cell.get("tooltip") == "Run checks in parallel" and "noLabel=1" in cell.find("mxCell").get("style")
